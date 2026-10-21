"""Создать отчёты DOCX из титульника пользователя, экранных форм и исходного кода.

Запуск в среде с python-docx, lxml и Pillow:
python tools/export_bot_reports.py --template /path/to/title.docx
PDF и изображения страниц создаются отдельно штатным render_docx.py.
"""
import argparse
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re
import textwrap
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
COMMON = ("telegram_labs/schedule.py", "telegram_labs/storage.py", "telegram_labs/app.py", "telegram_labs/launch.py")
TITLES = {
    5: "Разработка простого бота для Telegram с использованием языка Python",
    6: "Разработка бота на основе конечного автомата для Telegram с использованием языка Python",
}


def paragraph(doc, text="", size=12, bold=False, font="Times New Roman", keep=False):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_after = Pt(7)
    f.line_spacing = 1.15
    f.keep_with_next = keep
    f.widow_control = True
    run = p.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def heading(doc, text, page=False):
    p = paragraph(doc, text, size=14, bold=True, keep=True)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.page_break_before = page
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    p._p.get_or_add_pPr().append(outline)
    return p


def figure(doc, path, caption, width=6.0):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(width))
    paragraph(doc, caption, size=10)


def state_table(doc):
    rows = [
        ("Состояние", "Ввод и действие", "Следующий шаг"),
        ("SCHEDULE", "Текст или TXT; проверка расписания", "REMINDER_TIME"),
        ("REMINDER_TIME", "Время ЧЧ:ММ; пример одного дня", "CONFIRM"),
        ("CONFIRM", "Сохранить или отменить черновик", "END"),
    ]
    table = doc.add_table(rows=0, cols=3)
    table.autofit = False
    widths = (1.52, 3.13, 1.52)
    for i, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = 1
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                node = OxmlElement("w:" + edge)
                node.set(qn("w:w"), "90")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            tcpr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                node = OxmlElement("w:" + edge)
                for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    node.set(qn("w:" + key), val)
                borders.append(node)
            tcpr.append(borders)
            if i == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E9EDF1")
                tcpr.append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            r = p.add_run(value)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.bold = i == 0
            r.font.color.rgb = RGBColor(0, 0, 0)
        if i == 0:
            header = OxmlElement("w:tblHeader")
            table.rows[-1]._tr.get_or_add_trPr().append(header)


def append_code(doc, lab):
    lines = []
    for name in (f"lab_{lab:02}/main.py", *COMMON):
        lines += ["", "# Файл " + name, ""]
        for raw in (ROOT / name).read_text(encoding="utf-8").splitlines():
            # Unicode escapes в строках сохраняют смысл кода и печатаются без цветного emoji-шрифта.
            raw = "".join((f"\\U{ord(c):08x}" if ord(c) > 0xFFFF else
                           f"\\u{ord(c):04x}" if 0x2600 <= ord(c) <= 0x27BF or ord(c) == 0xFE0F else c)
                          for c in raw)
            lines.extend(textwrap.wrap(raw, width=84, replace_whitespace=False,
                                      drop_whitespace=False, subsequent_indent="    ") or [""])
    size = 60
    pages = (len(lines) + size - 1) // size
    balanced = (len(lines) + pages - 1) // pages
    for offset in range(0, len(lines), balanced):
        heading(doc, "Текст программы" if offset == 0 else "Текст программы продолжение", page=True)
        p = paragraph(doc, "\n".join(lines[offset:offset + balanced]), size=9, font="Courier New")
        p.paragraph_format.line_spacing = Pt(10.8)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.widow_control = False
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        p._p.get_or_add_pPr().append(snap)


def merge_preserving_title(body_doc, template, lab):
    stream = BytesIO()
    body_doc.save(stream)
    with ZipFile(template) as original, ZipFile(stream) as added:
        xml = etree.fromstring(original.read("word/document.xml"))
        for p in xml.xpath("//w:body/w:p", namespaces=NS):
            text = "".join(p.xpath(".//w:t/text()", namespaces=NS))
            replacement = None
            if text == "Лабораторная работа №1":
                replacement = f"Лабораторная работа №{lab}"
            if text == "«Разработка программы для решения биквадратного уравнения»":
                replacement = "«" + TITLES[lab] + "»"
            if replacement:
                nodes = p.xpath(".//w:t", namespaces=NS)
                nodes[0].text = replacement
                for node in nodes[1:]:
                    node.text = ""
        # Проверяющий на предоставленном титульнике уже указан верно.
        title_text = "".join(xml.xpath("//w:t/text()", namespaces=NS))
        assert "Гапанюк Ю.Е." in title_text
        body = xml.find("w:body", NS)
        appendix = etree.fromstring(added.read("word/document.xml")).find("w:body", NS)
        rels = etree.fromstring(original.read("word/_rels/document.xml.rels"))
        source_rels = etree.fromstring(added.read("word/_rels/document.xml.rels"))
        mapping, media = {}, {}
        for i, rel in enumerate(source_rels, 100):
            if not rel.get("Type").endswith("/image"):
                continue
            new_id = f"rIdBot{i}"
            target = f"media/bot{lab}_{Path(rel.get('Target')).name}"
            mapping[rel.get("Id")] = new_id
            etree.SubElement(rels, "{" + REL + "}Relationship", Id=new_id,
                             Type=rel.get("Type"), Target=target)
            media["word/" + target] = added.read("word/" + rel.get("Target"))
        for child in appendix:
            if child.tag == qn("w:sectPr"):
                continue
            for node in child.iter():
                for attr in (qn("r:embed"), qn("r:link")):
                    if node.get(attr) in mapping:
                        node.set(attr, mapping[node.get(attr)])
            body.insert(len(body) - 1, deepcopy(child))
        for i, node in enumerate(xml.xpath('//*[local-name()="docPr"]'), 1):
            if i > 1:
                node.set("id", str(100 + i))
        content_types = etree.fromstring(original.read("[Content_Types].xml"))
        if not any(n.get("Extension") == "png" for n in content_types):
            etree.SubElement(content_types, "{" + CT + "}Default", Extension="png", ContentType="image/png")
        modified = {
            "word/document.xml": etree.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True),
            "word/_rels/document.xml.rels": etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True),
            "[Content_Types].xml": etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone=True),
        }
        output = ROOT / f"lab_{lab:02}/report.docx"
        with ZipFile(output, "w", ZIP_DEFLATED) as target:
            for item in original.infolist():
                target.writestr(deepcopy(item), modified.get(item.filename, original.read(item.filename)))
            for name, raw in media.items():
                target.writestr(name, raw)
        with ZipFile(output) as final:
            for name in original.namelist():
                if name not in modified:
                    assert final.read(name) == original.read(name), name
    return output


def build(lab, template):
    doc = Document()
    doc.add_page_break()
    heading(doc, "Цель и задание")
    paragraph(doc, "Цель работы — изучить разработку ботов для Telegram на языке Python.")
    paragraph(doc, "Разработать простого Telegram-бота, использующего кнопки."
              if lab == 5 else "Разработать Telegram-бота, реализующего конечный автомат из трёх состояний.")
    paragraph(doc, "Реализован учебный ежедневник для группы ИУ5-32Б. Бот показывает занятия на сегодня, "
              "завтра и неделю, принимает новое расписание и отправляет занятия ежедневно в выбранное время.")
    heading(doc, "Описание реализации")
    paragraph(doc, "Использованы Python 3.10 или новее, python-telegram-bot 22.8, python-dotenv 1.2.1 и SQLite. "
              "Модули schedule.py, storage.py и app.py содержат общую логику двух лабораторных. "
              "У каждого личного чата отдельные расписание, время рассылки и отметка последней отправки.")
    paragraph(doc, "Меню создаётся ReplyKeyboardMarkup. Кнопки «Сегодня», «Завтра» и «Неделя» вызывают "
              "соответствующие обработчики; команды /today и текст «расписание» также возвращают занятия. "
              "После /start готовое расписание подключается автоматически, ежедневная отправка включается на 08:00.")
    paragraph(doc, "Учебные недели чередуются относительно известного понедельника числителя — 26 октября 2026 года. "
              "Неделя 19–25 октября является знаменателем. В исходном расписании 25 записей: 18 занятий "
              "на неделе ЗН и 16 на неделе ЧС. Пятница — день самостоятельной работы, воскресенье свободно.")
    paragraph(doc, "Для ответа по запросу используется дата отправки сообщения в Europe/Moscow; "
              "задержка обработки через полночь не меняет день ответа. JobQueue каждые 30 секунд проверяет "
              "время ежедневной отправки и вычисляет расписание на текущую дату. SQLite сохраняет подтверждённые "
              "данные и дату отправки после перезапуска.")
    if lab == 5:
        paragraph(doc, "Простой бот различает нажатия кнопок и данные расписания без конечного автомата. "
                  "Текст или TXT проверяется и сохраняется сразу. При ошибке старые данные остаются без изменений. "
                  "Время меняется через /notify ЧЧ:ММ, отправка выключается командой /off.")
    else:
        heading(doc, "Конечный автомат", page=True)
        paragraph(doc, "Изменение расписания запускается командой /upload. ConversationHandler содержит "
                  "ровно три состояния; до подтверждения данные находятся в context.user_data.")
        state_table(doc)
        paragraph(doc, "Ошибка формата оставляет пользователя в текущем состоянии. /cancel и /start завершают "
                  "диалог и удаляют черновик. END означает выход из автомата, а главное меню работает вне диалога. "
                  "Обновления обрабатываются последовательно через concurrent_updates(False).")
        paragraph(doc, "Расписание принимается текстом или TXT в UTF-8. Проверяются день, время, тип недели "
                  "и пересечения занятий. На шаге времени запись 8:00 нормализуется в 08:00. "
                  "При подтверждении показывается пример одного дня, затем расписание сохраняется в SQLite.")
    heading(doc, "Порядок запуска")
    paragraph(doc, "Из корня репозитория создать виртуальное окружение и установить requirements.txt. "
              "Скрипт запуска запрашивает токен скрытым вводом и сохраняет только в локальном .env, исключённом из Git.")
    for command in ("python3 -m venv .venv", "source .venv/bin/activate",
                    "python -m pip install -r requirements.txt", f"python telegram_labs/launch.py --lab {lab}"):
        paragraph(doc, command, size=9.5, font="Courier New")
    paragraph(doc, "В Telegram открыть @velisbro_study_diary_bot и отправить /start. "
              "Одновременно с одним токеном запускается только одна версия. Для ежедневной отправки нужны "
              "интернет и постоянно работающая программа; спящий компьютер сообщения не отправляет.")
    heading(doc, "Результаты выполнения", page=True)
    if lab == 5:
        paragraph(doc, "Экранные формы получены в реальном Telegram. На рисунках показаны текстовый ответ "
                  "с датой и типом недели, а также кнопка меню. Показан общий интерфейс модуля app.py, "
                  "который используют обе версии бота.")
        figure(doc, ROOT / "lab_05/results/telegram_today.png", "Рисунок 1 — Ответ на запрос текущего дня")
        figure(doc, ROOT / "lab_05/results/telegram_button.png", "Рисунок 2 — Кнопка Сегодня в клавиатуре бота")
    else:
        paragraph(doc, "Экранные формы получены в реальном Telegram при повторной загрузке того же расписания. "
                  "После демонстрации сохранено время 08:00 по Москве.")
        figure(doc, ROOT / "lab_06/results/telegram_fsm1.png", "Рисунок 1 — Состояние SCHEDULE и запрос расписания")
        figure(doc, ROOT / "lab_06/results/telegram_fsm2.png", "Рисунок 2 — Состояние REMINDER_TIME и ввод времени")
        heading(doc, "Подтверждение и результат", page=True)
        figure(doc, ROOT / "lab_06/results/telegram_fsm3.png", "Рисунок 3 — Состояние CONFIRM и проверка данных")
        figure(doc, ROOT / "lab_06/results/telegram_saved.png", "Рисунок 4 — Сохранение расписания и возврат в меню")
    heading(doc, "Проверка и выводы", page=True)
    count = 12 if lab == 5 else 6
    paragraph(doc, f"Для ЛР {lab} выполнено {count} проверок pytest; общий набор репозитория содержит 66 тестов, "
              "все завершились успешно. Журналы находятся в results/tests.txt и validation/tests.txt. "
              "Локальный транспорт подменяет Bot API, сохраняя настоящий диспетчер python-telegram-bot.")
    paragraph(doc, "Проверены импорт, ЧС/ЗН, сохранение, изоляция пользователей, отмена, время 8:00, "
              "повтор ответа после временной ошибки сети и запрос, задержанный через полночь. "
              "Реальная работа меню дополнительно проверена в Telegram и подтверждена экранными формами.")
    paragraph(doc, "Реализован ежедневник с кнопками и текстовой выдачей расписания. "
              + ("Кнопки вызывают обработчики запросов, а загрузка проверяет данные перед сохранением."
                 if lab == 5 else "Пошаговый ввод реализован конечным автоматом из трёх состояний с отменой и подтверждением."))
    paragraph(doc, "Пропущенные при простое дни не рассылаются, после запуска возможна отправка за текущий день. "
              "Каникулы и переносы занятий требуют обновления расписания. При аварии после отправки, "
              "но до записи отметки в SQLite возможен повтор сообщения.")
    heading(doc, "Источники")
    for url in (f"https://github.com/ugapanyuk/BKIT_2021/wiki/lab_bot{lab - 4}",
                "https://github.com/ugapanyuk/courses_content/wiki/COURSE_PCPL_MAIN",
                "https://docs.python-telegram-bot.org/en/v22.8/", "https://core.telegram.org/bots/api"):
        paragraph(doc, url, size=9.5)
    paragraph(doc, "Далее приведены полные исходные тексты основного и общих модулей, включая скрипт запуска. "
              "Эмодзи в печатном листинге представлены эквивалентными Unicode escape последовательностями.", size=10)
    append_code(doc, lab)
    return merge_preserving_title(doc, template, lab)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    args = parser.parse_args()
    for lab in (5, 6):
        print(build(lab, args.template))


if __name__ == "__main__":
    main()
